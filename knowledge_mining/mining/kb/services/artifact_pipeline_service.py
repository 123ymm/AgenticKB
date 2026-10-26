"""Artifact table pipeline: durable inputs/outputs, existing queue and object store."""
from __future__ import annotations

import asyncio
import hashlib
import json
import logging
import uuid
from contextlib import suppress

from knowledge_mining.mining.contracts.file_management import StorageObjectRecord
from knowledge_mining.mining.contracts.storage.types import ObjectLocation, PutOptions
from knowledge_mining.mining.file_management.repositories_pg import PgStorageObjectRepository
from knowledge_mining.mining.infra.object_store.keys import build_object_key
from knowledge_mining.mining.kb.artifact_pipeline_repository import ArtifactPipelineRepository
from knowledge_mining.mining.kb.db import KbDB
from knowledge_mining.mining.kb.db_artifacts import ArtifactDB
from knowledge_mining.mining.kb.services.artifact_service import (
    ArtifactService, ArtifactNotFound, InvalidOperation, ROLE_DESIGN,
)
from knowledge_mining.mining.kb.services.kb_service import KbService, NotFound
from knowledge_mining.mining.workflow.artifact_tables import (
    COLUMNS, PROFILE, OPERATOR, extract_rows, extraction_manifest, table_handler_registry,
)
from knowledge_mining.mining.workflow.manifest import value_hash
from knowledge_mining.mining.workflow.artifact_table_ops import (
    CATALOG, materialize, selection, transform, validate_columns, cast,
)

logger = logging.getLogger(__name__)
MAX_SNAPSHOT_BYTES = 64 * 1024 * 1024


class ArtifactPipelineService:
    def __init__(self, *, repository, artifacts, kbs, kb_db, storage, objects, bucket, provider):
        self.repo = repository
        self.artifacts = artifacts
        self.kbs = kbs
        self.kb_db = kb_db
        self.storage = storage
        self.objects = objects
        self.bucket = bucket
        self.provider = provider

    async def authorize(self, user: dict, artifact_id: str, kb_id: str | None = None):
        artifact, _ = await self.artifacts._assert_role(user, artifact_id, *ROLE_DESIGN)
        if kb_id:
            try:
                await self.kbs._assert_read(kb_id, str(user["id"]))
                kb = await self.kb_db.get_kb(kb_id)
            except NotFound:
                raise ArtifactNotFound(artifact_id) from None
            if not kb or kb["domain"] != artifact["domain_id"]:
                raise ArtifactNotFound(artifact_id)
        return artifact

    async def submit(self, user: dict, artifact_id: str, kb_id: str, document_ids: list[str]):
        artifact = await self.authorize(user, artifact_id, kb_id)
        ids = list(dict.fromkeys(document_ids))
        if not ids or len(ids) > 2000:
            raise InvalidOperation("请选择 1–2000 个文件")
        sources = await self.repo.freeze_sources(kb_id, ids)
        if {s["id"] for s in sources} != set(ids):
            raise InvalidOperation("所选文件已移动或删除，请重新选择来源")
        return await self.repo.enqueue(artifact, str(user["id"]), kb_id, extraction_manifest(sources))

    @staticmethod
    def summary(run):
        return {k: run.get(k) for k in ("id", "status", "started_at", "finished_at", "error_summary")} | {
            "output": run["metadata_json"].get("output", {}),
            **{k: run["metadata_json"].get(k) for k in (
                "parent_run_id", "step_index", "operator", "progress")},
        }

    async def submit_step(self, user, artifact_id, parent_id, operator, parameters):
        if operator not in dict(CATALOG): raise InvalidOperation("未知表格算子")
        parent = await self.read_run(user, artifact_id, parent_id)
        if parent["status"] != "completed" or not parent["metadata_json"].get("output"):
            raise InvalidOperation("请先完成输入步骤")
        # Object references and source lineage come only from authorized persisted runs.
        params = dict(parameters)
        params.pop("input_object_id", None)
        params.pop("sources", None)
        params.pop("lookup_object_id", None)
        if operator == "join":
            lookup = await self.lookup(user, artifact_id, params.get("lookup_id", ""))
            params["lookup_object_id"] = lookup["storage_object_id"]
            params["lookup_columns"] = lookup["columns"]
            params["lookup_filename"] = lookup["filename"]
            params["lookup_sha256"] = lookup["sha256"]
        manifest = parent["workflow_manifest_json"]
        params["sources"] = manifest["nodes"][0]["params"]["sources"]
        params["input_object_id"] = parent["metadata_json"]["output"]["storage_object_id"]
        snapshot = await self.read_json(params["input_object_id"])
        if operator == "llm_columns":
            selected = await asyncio.to_thread(selection, snapshot, params)
            retry = bool(params.get("retry_failed"))
            if retry:
                old = manifest["nodes"][0]
                if old["operatorType"] != "llm_columns": raise InvalidOperation("只能重试大模型步骤")
                # Keep original prompt/schema; only failed rows will run again.
                params = dict(old["params"]) | {"input_object_id": params["input_object_id"], "retry_failed": True, "preview": False}
            else:
                validate_columns(params.get("columns", []), {c["key"] for c in selected["columns"]})
                if any(r.get("_llm", {}).get("status") == "failed" for r in selected["rows"]):
                    raise InvalidOperation("输入有模型失败行，请先重试该步骤或筛选移除失败行")
            if not params.get("prompt", "").strip(): raise InvalidOperation("请填写大模型加工要求")
            if not set(params.get("input_columns", [])) <= {c["key"] for c in selected["columns"]} or not params.get("input_columns"):
                raise InvalidOperation("请选择大模型输入列")
        else:
            # Validate input selection cheaply; full transforms run only in the existing async queue.
            try:
                await asyncio.to_thread(selection, snapshot, params)
            except (ValueError, TypeError, KeyError) as exc:
                raise InvalidOperation(f"算子参数无效：{exc}") from None
        artifact = await self.authorize(user, artifact_id)
        return await self.repo.enqueue(artifact, str(user["id"]), parent["metadata_json"]["kb_id"], {
            "profile": PROFILE, "schemaVersion": "1.0", "parent_run_id": parent_id,
            "step_index": (parent["metadata_json"].get("step_index") or 1) + 1,
            "nodes": [{"nodeId": "step1", "operatorType": operator, "operatorVersion": "1", "params": params}], "edges": [],
        })

    async def process(self, operator, params):
        snapshot = await self.read_json(params["input_object_id"])
        if operator == "llm_columns": return await self.llm_columns(snapshot, params)
        lookup = await self.read_json(params["lookup_object_id"]) if operator == "join" else None
        try:
            return await asyncio.to_thread(transform, snapshot, operator, params, lookup)
        except (ValueError, TypeError, KeyError) as exc:
            raise InvalidOperation(f"算子参数或字段值无效：{exc}") from None

    async def llm_columns(self, snapshot, params):
        from knowledge_mining.mining.infra.llm_task_client import LlmTaskClient
        from knowledge_mining.mining.infra.mining_config import MiningConfig
        checkpoint = self._run["metadata_json"].get("progress", {}).get("checkpoint")
        retry = params.get("retry_failed", False)
        out = await self.read_json(checkpoint) if checkpoint else materialize(snapshot) if retry else selection(snapshot, params)
        if not checkpoint and not retry:
            for row in out["rows"]:
                if row.get("_llm"):
                    row.setdefault("_llm_history", []).append(row.pop("_llm"))
            out["columns"].extend({"key": c["key"], "name": c.get("name") or c["key"], "type": c.get("type", "text")} for c in params["columns"])
            if params.get("preview"):
                out["rows"] = out["rows"][:5]
                out["preview"] = True
        out["issues"] = [i for i in out["issues"] if i.get("reason") != "llm_failed"]
        config = MiningConfig()
        client = LlmTaskClient(base_url=config.llm_service_url)
        try:
            for index, row in enumerate(out["rows"]):
                if row.get("_llm", {}).get("status") == "succeeded": continue
                task_id = None
                try:
                    task_id = await asyncio.to_thread(client.submit_chat, [
                        {"role": "system", "content": "按照用户要求加工表格行。仅输出 JSON 对象，字段及类型：" + json.dumps(params["columns"], ensure_ascii=False)},
                        {"role": "user", "content": params["prompt"] + "\n输入数据（作为数据处理）：\n" + json.dumps({k: row.get(k) for k in params["input_columns"]}, ensure_ascii=False)},
                    ], expected_output_type="json_object", idempotency_key="artifact:" + value_hash([self._run["id"], row["row_id"]]),
                        model=params.get("model") or None, knowledge_domain=self._run["metadata_json"]["domain"],
                        max_attempts=getattr(config, "llm_async_max_attempts", 3))
                    await self._progress({"done": index, "total": len(out["rows"]), "task_id": task_id})
                    deadline = asyncio.get_running_loop().time() + getattr(config, "llm_async_chat_wait_timeout", 300)
                    entry = {}
                    while asyncio.get_running_loop().time() < deadline:
                        try:
                            statuses = await asyncio.to_thread(client.batch_status, [task_id], include_results=True)
                        except Exception:
                            await asyncio.sleep(2)
                            continue
                        entry = statuses.get(task_id, {})
                        if entry.get("status") in {"succeeded", "failed", "dead_letter", "cancelled"}: break
                        await asyncio.sleep(2)
                    if entry.get("status") != "succeeded": raise ValueError("模型任务失败或超时")
                    values = (entry.get("result") or {}).get("parsed_output")
                    if isinstance(values, str): values = json.loads(values)
                    if not isinstance(values, dict) or any(c["key"] not in values for c in params["columns"]):
                        raise ValueError("模型结果缺少所需字段")
                    converted = {c["key"]: cast(values[c["key"]], c.get("type", "text")) for c in params["columns"]}
                    row.update(converted)
                    row["_llm"] = {"status": "succeeded", "task_id": task_id, "run_id": self._run["id"], "model": params.get("model") or "default"}
                except asyncio.CancelledError: raise
                except Exception:
                    logger.exception("Workbench LLM task failed: %s", task_id)
                    row["_llm"] = {"status": "failed", "task_id": task_id}
                progress = {"done": index + 1, "total": len(out["rows"]), "task_id": task_id}
                if (index + 1) % 10 == 0 or index == len(out["rows"]) - 1:
                    saved = await self.write_snapshot(out)
                    progress["checkpoint"] = saved["storage_object_id"]
                await self._progress(progress)
        finally:
            await asyncio.to_thread(client.close)
        out["issues"].extend({"row_id": r["row_id"], "reason": "llm_failed", "message": "模型加工失败，可重试失败行", "task_id": r.get("_llm", {}).get("task_id")}
                             for r in out["rows"] if r.get("_llm", {}).get("status") == "failed")
        return out

    async def _progress(self, values):
        self._run["metadata_json"].setdefault("progress", {}).update(values)
        if not await self.repo.progress(self._run["id"], self._worker, self._run["metadata_json"]["progress"]):
            raise asyncio.CancelledError()

    async def upload_lookup(self, user, artifact_id, filename, content):
        import io
        from openpyxl import load_workbook
        from knowledge_mining.mining.kb.services.artifact_import_service import _check_zip_bomb, _dedupe_headers
        from knowledge_mining.mining.workflow.repositories.global_workflow_repository import GlobalWorkflowRepository
        await self.authorize(user, artifact_id)
        if not filename.lower().endswith(".xlsx"): raise InvalidOperation("对照表请选择 .xlsx 文件")
        _check_zip_bomb(content)
        def parse():
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                sheet = workbook.worksheets[0]
                values = sheet.iter_rows(values_only=True)
                header = next(values, ())
                if not header or len(header) > 200: raise InvalidOperation("对照表需首行为表头，最多 200 列")
                names = _dedupe_headers(header)
                used = set()
                for i, name in enumerate(names):
                    suffix = 2
                    unique = name
                    while unique in used:
                        unique = f"{name}_{suffix}"
                        suffix += 1
                    names[i] = unique
                    used.add(unique)
                from knowledge_mining.mining.workflow.artifact_table_ops import RESERVED
                if set(names) & RESERVED: raise InvalidOperation("表头包含保留字段，请修改表头")
                rows = []
                for row in values:
                    if not any(v is not None for v in row): continue
                    if any(isinstance(v, str) and v.startswith("=") for v in row):
                        raise InvalidOperation("对照表含公式，请粘贴为值后上传")
                    rows.append({"row_id": str(len(rows)), **{k: cast(v, "text") for k, v in zip(names, row)}})
                    if len(rows) > 100000: raise InvalidOperation("对照表最多 100000 行")
                return {"schema_version": 1, "columns": [{"key": k, "name": k, "type": "text"} for k in names], "rows": rows, "issues": [], "documents": {}}
            finally: workbook.close()
        try: snapshot = await asyncio.to_thread(parse)
        except InvalidOperation: raise
        except Exception: raise InvalidOperation("无法读取对照表，请检查 Excel 格式") from None
        output = await self.write_snapshot(snapshot)
        lookup_id = uuid.uuid4().hex
        await GlobalWorkflowRepository(self.repo.pool).insert_workflow({
            "id": lookup_id, "name": f"artifact-lookup-{lookup_id}", "draft_graph_json": {},
            "created_by": str(user["id"]), "metadata_json": {"profile": PROFILE, "kind": "lookup",
                "artifact_id": artifact_id, "user_id": str(user["id"]), "filename": filename,
                "output": output, "columns": snapshot["columns"]},
        })
        return {"lookup_id": lookup_id, "filename": filename, "columns": snapshot["columns"], "row_count": len(snapshot["rows"])}

    async def lookup(self, user, artifact_id, lookup_id):
        await self.authorize(user, artifact_id)
        async with self.repo.pool.connection() as conn:
            cursor = await conn.execute("SELECT metadata_json FROM mining_workflows WHERE id=%s AND metadata_json->>'artifact_id'=%s AND metadata_json->>'kind'='lookup' AND metadata_json->>'profile'=%s",
                                        (lookup_id, artifact_id, PROFILE))
            record = await cursor.fetchone()
            if not record: raise ArtifactNotFound(lookup_id)
            metadata = record["metadata_json"]
            return metadata["output"] | {"columns": metadata["columns"], "filename": metadata["filename"]}

    async def list_runs(self, user: dict, artifact_id: str):
        await self.authorize(user, artifact_id)
        runs = await self.repo.list_runs(artifact_id)
        # A private source KB must not become readable through an artifact membership.
        visible = []
        for run in runs:
            try:
                await self.authorize(user, artifact_id, run["metadata_json"]["kb_id"])
                visible.append(self.summary(run))
            except ArtifactNotFound:
                continue
        return visible

    async def read_run(self, user: dict, artifact_id: str, run_id: str):
        await self.authorize(user, artifact_id)
        run = await self.repo.get_run(artifact_id, run_id)
        if not run:
            raise ArtifactNotFound(run_id)
        await self.authorize(user, artifact_id, run["metadata_json"]["kb_id"])
        return run

    async def page(self, user: dict, artifact_id: str, run_id: str, page: int, size: int, filters=None, keep_columns=None):
        run = await self.read_run(user, artifact_id, run_id)
        result = self.summary(run)
        sources = run["workflow_manifest_json"]["nodes"][0]["params"]["sources"]
        result.update({"sources": sources, "columns": COLUMNS, "rows": [], "total": 0, "issues": []})
        result["kb_id"] = run["metadata_json"]["kb_id"]
        result["publications"] = []
        if run.get("workflow_id"):
            from knowledge_mining.mining.kb.services.artifact_service import ArtifactError
            for published in await self.repo.publications(run["workflow_id"]):
                try: await self.artifacts.get_artifact(user=user, artifact_id=published["artifact_id"])
                except ArtifactError: continue
                result["publications"].append(published)
        result["parameters"] = {k: v for k, v in run["workflow_manifest_json"]["nodes"][0]["params"].items()
                                if k not in {"sources", "input_object_id", "lookup_object_id"}}
        output_id = run["metadata_json"].get("output", {}).get("storage_object_id")
        if output_id:
            snapshot = selection(await self.read_json(output_id), {"filters": filters or [], "keep_columns": keep_columns})
            rows = snapshot["rows"][(page - 1) * size:page * size]
            result.update({
                "total": len(snapshot["rows"]), "issues": snapshot["issues"],
                "rows": rows, "columns": snapshot["columns"], "preview": snapshot.get("preview", False),
            })
        return result

    async def read_json(self, object_id: str) -> dict:
        obj = await self.objects.get(object_id)
        if not obj or obj.state != "AVAILABLE":
            raise InvalidOperation("快照文件不可用，请检查对象存储")
        if obj.size > MAX_SNAPSHOT_BYTES:
            raise InvalidOperation("单个快照超过 64MB，请缩小来源范围")
        raw = bytearray()
        async for chunk in self.storage.get_stream(ObjectLocation(obj.bucket, obj.object_key, obj.object_version_id)):
            raw.extend(chunk)
            if len(raw) > MAX_SNAPSHOT_BYTES:
                raise InvalidOperation("单个快照超过 64MB，请缩小来源范围")
        if hashlib.sha256(raw).hexdigest() != obj.sha256:
            raise InvalidOperation("快照内容校验失败")
        return json.loads(raw)

    async def write_snapshot(self, snapshot: dict) -> dict:
        raw = json.dumps(snapshot, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode()
        if len(raw) > MAX_SNAPSHOT_BYTES:
            raise InvalidOperation("结果超过 64MB，请缩小来源范围；本次未截断保存")
        object_id = uuid.uuid4().hex
        # binary bucket has no staging expiration; immutable key contains no filenames/credentials.
        location = ObjectLocation(self.bucket, build_object_key(
            "binary_asset", hashlib.sha256(raw).hexdigest(), prefix="artifact-tables/v1",
        ))

        async def chunks():
            yield raw

        result = await self.storage.put_stream(location, chunks(), PutOptions(
            artifact_class="binary_asset", mime="application/json", content_length=len(raw),
            expected_sha256=hashlib.sha256(raw).hexdigest(),
        ))
        record = await self.objects.register(StorageObjectRecord(
            id=object_id, provider=self.provider, bucket=self.bucket, object_key=location.object_key,
            object_version_id=result.version_id, sha256=result.sha256, size=result.size,
            mime="application/json", artifact_class="binary_asset", state="AVAILABLE", etag=result.etag,
        ))
        return {"storage_object_id": record.id, "sha256": record.sha256,
                "row_count": len(snapshot["rows"]), "issue_count": len(snapshot["issues"])}

    async def extract(self, sources: list[dict]) -> dict:
        rows, documents, issues = [], {}, []
        byte_count = 0
        for source in sources:
            issue = {"document_id": source["id"], "file_name": source.get("document_name") or source["id"]}
            if not source.get("snapshot_id"):
                issues.append(issue | {"reason": "not_extracted", "message": "当前文件版本尚无可用抽取快照"})
                continue
            try:
                tables, cells = await self.repo.table_data(source["snapshot_id"])
                ir = await self.read_json(source["parse_ir_storage_object_id"])
                # IR knows about a table but projections are absent: never claim this file has no tables.
                ir_tables = [a for a in (ir.get("structured_assets") or {}).values() if "table_id" in a]
                if not tables:
                    issues.append(issue | {
                        "reason": "projection_missing" if ir_tables else "no_tables",
                        "message": "表格尚未完成结构化入库，请完成该文件挖掘" if ir_tables else "该文件未抽取到表格",
                    })
                    continue
                if any(not any(c["table_ref"] == t["table_ref"] for c in cells)
                       and (t.get("row_count") or 0) > 0 for t in tables):
                    raise InvalidOperation("表格单元格数据不完整，请重新挖掘该文件")
                extracted, text = extract_rows(source, tables, cells, ir)
            except Exception:
                logger.exception("Artifact table extraction failed for document %s", source["id"])
                issues.append(issue | {"reason": "read_failed", "message": "表格或正文读取失败，请检查挖掘记录及对象存储"})
                continue
            byte_count += len(json.dumps(extracted, ensure_ascii=False).encode()) + len(text.encode())
            if byte_count > MAX_SNAPSHOT_BYTES:
                raise InvalidOperation("结果超过 64MB，请缩小来源范围；本次未截断保存")
            rows.extend(extracted)
            documents[source["id"]] = text
        return {"schema_version": 1, "columns": COLUMNS, "rows": rows, "documents": documents, "issues": issues}

    async def execute(self, candidate: dict):
        run_id = str(candidate["id"])
        artifact_id = candidate["metadata_json"]["artifact_id"]
        run = await self.repo.get_run(artifact_id, run_id)
        if not run:
            return {"status": "claimed_elsewhere"}
        worker_id = f"artifact-{uuid.uuid4().hex}"
        operator = (run["workflow_manifest_json"].get("nodes") or [{}])[0].get("operatorType", OPERATOR)
        event_id = await self.repo.claim(run_id, worker_id, operator)
        if event_id is None:
            return {"status": "claimed_elsewhere"}
        current = asyncio.current_task()

        async def renew():
            while True:
                await asyncio.sleep(30)
                try:
                    owned = await self.repo.heartbeat(run_id, worker_id)
                except Exception:
                    owned = False
                if not owned:
                    current.cancel()
                    return

        heartbeat = asyncio.create_task(renew())
        try:
            metadata = run["metadata_json"]
            user = await self.kb_db.get_user(metadata["user_id"])
            if not user or user.get("deleted_at") or user.get("status") == "disabled":
                raise InvalidOperation("操作者已失效")
            await self.authorize(user, artifact_id, metadata["kb_id"])
            manifest = run["workflow_manifest_json"]
            if value_hash(manifest) != run["workflow_graph_hash"] or manifest.get("profile") != PROFILE:
                raise InvalidOperation("加工配置校验失败")
            node = manifest["nodes"][0]
            self._run, self._worker = run, worker_id
            handler = table_handler_registry().resolve(node["operatorType"], node["operatorVersion"])
            snapshot = await handler(self, node["params"], None)
            output = await self.write_snapshot(snapshot)
            await self.repo.finish(run_id, event_id, worker_id, output=output)
        except asyncio.CancelledError:
            # Leave the lease/recovery path in charge; never fail a newer owner's attempt.
            return {"status": "claimed_elsewhere"}
        except Exception as exc:
            logger.exception("Artifact pipeline run %s failed", run_id)
            await self.repo.finish(run_id, event_id, worker_id, output={},
                                   error=str(exc) if isinstance(exc, InvalidOperation) else "加工失败，请检查服务日志、来源权限或缩小选择范围后重试")
        finally:
            heartbeat.cancel()
            with suppress(asyncio.CancelledError):
                await heartbeat
        return {"status": "done"}


async def service_for(app, domain: str) -> ArtifactPipelineService:
    pool = app.state.pg_pool
    runtime_pool = await app.state.domain_pools.async_pool(domain)
    config = app.state.object_store_config
    return ArtifactPipelineService(
        repository=ArtifactPipelineRepository(pool, runtime_pool),
        artifacts=ArtifactService(ArtifactDB(pool)), kbs=KbService(KbDB(pool)), kb_db=KbDB(pool),
        storage=app.state.object_store, objects=PgStorageObjectRepository(pool),
        bucket=f"{config.bucket_prefix}binary", provider=config.provider,
    )
